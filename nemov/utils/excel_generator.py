"""
Excel Report Generator for Yalda Gym Application.
Produces presentation-grade, professional .xlsx workbooks with Persian RTL orientation,
executive KPI summary cards, styled tables, formatted numbers, and unlocked editable cells.
"""

from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter

from nemov.utils.jalali_date import get_today_shamsi, get_now_shamsi_for_filename
from nemov.utils.logger import log_info, log_error

class ExcelGenerator:
    RTL_READING_ORDER = 2  # Excel Right-to-Left text direction

    @staticmethod
    def _rtl_align(horizontal="center", vertical="center", **kwargs) -> Alignment:
        """Creates an openpyxl Alignment with strict Right-to-Left readingOrder=2."""
        return Alignment(
            horizontal=horizontal,
            vertical=vertical,
            readingOrder=ExcelGenerator.RTL_READING_ORDER,
            **kwargs
        )

    @staticmethod
    def _style_range(ws, cell_range, font=None, fill=None, border=None, alignment=None):
        """Helper to style all cells in a range (especially merged ranges)."""
        no_lock = Protection(locked=False)
        if alignment is not None and getattr(alignment, "readingOrder", None) != ExcelGenerator.RTL_READING_ORDER:
            alignment = Alignment(
                horizontal=alignment.horizontal,
                vertical=alignment.vertical,
                wrap_text=alignment.wrap_text,
                shrink_to_fit=alignment.shrink_to_fit,
                indent=alignment.indent,
                text_rotation=alignment.text_rotation,
                readingOrder=ExcelGenerator.RTL_READING_ORDER,
            )
        for row in ws[cell_range]:
            for cell in row:
                if font:
                    cell.font = font
                if fill:
                    cell.fill = fill
                if border:
                    cell.border = border
                if alignment:
                    cell.alignment = alignment
                cell.protection = no_lock

    @staticmethod
    def generate_financial_excel(
        output_filepath: str,
        summary: dict,
        transactions: list,
        filter_info: dict = None
    ) -> str:
        """
        Creates a beautifully styled, unlocked .xlsx workbook for financial statements.
        
        Args:
            output_filepath: Destination file path for .xlsx file.
            summary: Dict containing total_income, total_expense, net_profit, total_transactions.
            transactions: List of transaction dicts.
            filter_info: Dict with 'time_label' and 'trainer_label'.
            
        Returns:
            str: Path of generated file.
        """
        output_path = Path(output_filepath)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        title_font_family = "B Titr"
        font_family = "B Koodak"

        # Ensure default workbook style enforces B Koodak and Right-to-Left (readingOrder=2) for all cells (including empty/new cells)
        try:
            for style in wb._named_styles:
                if getattr(style, "name", "") == "Normal":
                    style.font = Font(name=font_family, size=11)
                    style.alignment = Alignment(readingOrder=ExcelGenerator.RTL_READING_ORDER)
        except Exception:
            pass

        ws = wb.active
        ws.title = "صورت وضعیت مالی"

        # 1. RTL Orientation & View Options
        ws.sheet_view.rightToLeft = True
        ws.views.sheetView[0].showGridLines = True
        ws.protection.sheet = False  # Completely editable / unlocked

        # Page Setup for Printing / Export (Landscape A4 fit to 1 page wide)
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

        # Reusable Typography & Styles
        border_thin = Side(style="thin", color="CBD5E1")
        border_double = Side(style="double", color="64748B")
        cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
        no_lock = Protection(locked=False)

        # ----------------------------------------------------
        # Row 1: Top Margin Spacer
        # ----------------------------------------------------
        ws.row_dimensions[1].height = 10

        # ----------------------------------------------------
        # Row 2: Title Banner (Merged A2:H2)
        # ----------------------------------------------------
        ws.merge_cells("A2:H2")
        title_font = Font(name=title_font_family, size=15, bold=True, color="FFFFFF")
        title_fill = PatternFill(start_color="881337", end_color="881337", fill_type="solid")  # Brand Deep Maroon
        title_align = ExcelGenerator._rtl_align(horizontal="center", vertical="center")
        title_cell = ws["A2"]
        from nemov.services.branding_service import BrandingService
        club_name = BrandingService.get_active_club_name()
        title_cell.value = f"صورت وضعیت مالی و تراز درآمد و هزینه‌های {club_name}"
        ExcelGenerator._style_range(ws, "A2:H2", font=title_font, fill=title_fill, alignment=title_align)
        ws.row_dimensions[2].height = 42

        # ----------------------------------------------------
        # Row 3: Sub-Header / Metadata Info Bar (Merged A3:H3)
        # ----------------------------------------------------
        ws.merge_cells("A3:H3")
        time_lbl = (filter_info or {}).get("time_label", "کل دوره")
        trainer_lbl = (filter_info or {}).get("trainer_label", "همه مربیان (کل باشگاه)")
        today_fa = get_today_shamsi()

        meta_font = Font(name=font_family, size=11, bold=True, color="334155")
        meta_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        meta_align = ExcelGenerator._rtl_align(horizontal="center", vertical="center")
        meta_border = Border(top=border_thin, bottom=border_thin, left=border_thin, right=border_thin)

        meta_cell = ws["A3"]
        meta_cell.value = f"\u200Fبازه زمانی: \u200F{time_lbl}\u200F   |   \u200Fمربی انتخابی: \u200F{trainer_lbl}\u200F   |   \u200Fتاریخ گزارش‌گیری: \u200F{today_fa}\u200F"
        ExcelGenerator._style_range(ws, "A3:H3", font=meta_font, fill=meta_fill, border=meta_border, alignment=meta_align)
        ws.row_dimensions[3].height = 26

        # ----------------------------------------------------
        # Row 4: Spacer
        # ----------------------------------------------------
        ws.row_dimensions[4].height = 12

        # ----------------------------------------------------
        # Rows 5 & 6: Executive KPI / Summary Stat Cards
        # Card 1 (A-B): مجموع درآمدها
        # Card 2 (C-D): مجموع هزینه‌ها
        # Card 3 (E-F): سود و تراز خالص
        # Card 4 (G-H): تعداد تراکنش‌ها
        # ----------------------------------------------------
        tot_inc = summary.get("total_income", 0.0)
        tot_exp = summary.get("total_expense", 0.0)
        profit = summary.get("net_profit", 0.0)
        tot_count = summary.get("total_transactions", len(transactions))

        cards = [
            {
                "cols": ("A", "B"),
                "title": "مجموع درآمدها (شهریه)",
                "value": f"\u200F\u200E{int(tot_inc):,}\u200F تومان",
                "bg": "ECFDF5",
                "title_color": "065F46",
                "val_color": "047857",
                "border_color": "A7F3D0"
            },
            {
                "cols": ("C", "D"),
                "title": "مجموع هزینه‌ها (مخارج)",
                "value": f"\u200F\u200E{int(tot_exp):,}\u200F تومان",
                "bg": "FEF2F2",
                "title_color": "991B1B",
                "val_color": "B91C1C",
                "border_color": "FECACA"
            },
            {
                "cols": ("E", "F"),
                "title": "سود و تراز خالص دوره",
                "value": f"\u200F\u200E{'+' if profit > 0 else ''}{int(profit):,}\u200F تومان",
                "bg": "EFF6FF" if profit >= 0 else "FFF1F2",
                "title_color": "1E40AF" if profit >= 0 else "991B1B",
                "val_color": "1D4ED8" if profit >= 0 else "DC2626",
                "border_color": "BFDBFE" if profit >= 0 else "FECACA"
            },
            {
                "cols": ("G", "H"),
                "title": "تعداد کل تراکنش‌ها",
                "value": f"\u200F\u200E{tot_count}\u200F فقره",
                "bg": "F8FAFC",
                "title_color": "475569",
                "val_color": "1E293B",
                "border_color": "E2E8F0"
            }
        ]

        ws.row_dimensions[5].height = 22
        ws.row_dimensions[6].height = 28

        for card in cards:
            c1, c2 = card["cols"]
            ws.merge_cells(f"{c1}5:{c2}5")
            ws.merge_cells(f"{c1}6:{c2}6")

            t_font = Font(name=title_font_family, size=10, bold=True, color=card["title_color"])
            v_font = Font(name=font_family, size=13, bold=True, color=card["val_color"])
            card_fill = PatternFill(start_color=card["bg"], end_color=card["bg"], fill_type="solid")
            b_side = Side(style="thin", color=card["border_color"])
            c_border = Border(left=b_side, right=b_side, top=b_side, bottom=b_side)
            c_align = ExcelGenerator._rtl_align(horizontal="center", vertical="center")

            ws[f"{c1}5"].value = card["title"]
            ws[f"{c1}6"].value = card["value"]

            ExcelGenerator._style_range(ws, f"{c1}5:{c2}5", font=t_font, fill=card_fill, border=c_border, alignment=c_align)
            ExcelGenerator._style_range(ws, f"{c1}6:{c2}6", font=v_font, fill=card_fill, border=c_border, alignment=c_align)

        # ----------------------------------------------------
        # Row 7: Spacer
        # ----------------------------------------------------
        ws.row_dimensions[7].height = 14

        # ----------------------------------------------------
        # Row 8: Table Headers
        # ----------------------------------------------------
        headers = [
            ("ردیف", "center"),
            ("نوع تراکنش", "center"),
            ("عنوان / طرف حساب", "right"),
            ("دسته‌بندی", "center"),
            ("تاریخ شمسی", "center"),
            ("مبلغ (تومان)", "center"),
            ("مربی مربوطه", "center"),
            ("توضیحات", "right"),
        ]

        header_fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")  # Crimson 800
        header_font = Font(name=title_font_family, size=10.5, bold=True, color="FFFFFF")
        ws.row_dimensions[8].height = 30

        for col_idx, (hdr_text, align) in enumerate(headers, start=1):
            cell = ws.cell(row=8, column=col_idx)
            cell.value = hdr_text
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = ExcelGenerator._rtl_align(horizontal="center", vertical="center")
            cell.border = cell_border
            cell.protection = no_lock

        # ----------------------------------------------------
        # Rows 9+: Data Rows
        # ----------------------------------------------------
        zebra_fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        zebra_fill_gray = PatternFill(start_color="F8FAFB", end_color="F8FAFB", fill_type="solid")
        font_regular = Font(name=font_family, size=11, bold=False, color="1E293B")
        font_bold = Font(name=font_family, size=11, bold=True, color="1E293B")

        current_row = 9
        for idx, t in enumerate(transactions, start=1):
            ws.row_dimensions[current_row].height = 25
            is_income = t.get("type") == "income"
            row_fill = zebra_fill_gray if idx % 2 == 0 else zebra_fill_white

            # 1. ردیف
            c_idx = ws.cell(row=current_row, column=1, value=idx)
            c_idx.font = font_regular
            c_idx.alignment = ExcelGenerator._rtl_align(horizontal="center", vertical="center")
            c_idx.fill = row_fill
            c_idx.border = cell_border
            c_idx.protection = no_lock

            # 2. نوع تراکنش
            type_label = "درآمد (شهریه)" if is_income else "هزینه باشگاه"
            c_type = ws.cell(row=current_row, column=2, value=type_label)
            c_type.font = Font(name=font_family, size=10.5, bold=True, color="047857" if is_income else "B91C1C")
            c_type.fill = PatternFill(start_color="ECFDF5" if is_income else "FEF2F2", end_color="ECFDF5" if is_income else "FEF2F2", fill_type="solid")
            c_type.alignment = ExcelGenerator._rtl_align(horizontal="center", vertical="center")
            c_type.border = cell_border
            c_type.protection = no_lock

            # 3. عنوان / طرف حساب
            c_title = ws.cell(row=current_row, column=3, value=t.get("title", "-") or "-")
            c_title.font = font_bold
            c_title.alignment = ExcelGenerator._rtl_align(horizontal="right", vertical="center", indent=1)
            c_title.fill = row_fill
            c_title.border = cell_border
            c_title.protection = no_lock

            # 4. دسته‌بندی
            c_cat = ws.cell(row=current_row, column=4, value=t.get("category", "-") or "-")
            c_cat.font = font_regular
            c_cat.alignment = ExcelGenerator._rtl_align(horizontal="center", vertical="center")
            c_cat.fill = row_fill
            c_cat.border = cell_border
            c_cat.protection = no_lock

            # 5. تاریخ شمسی
            c_date = ws.cell(row=current_row, column=5, value=t.get("date_shamsi", "-") or "-")
            c_date.font = font_regular
            c_date.alignment = ExcelGenerator._rtl_align(horizontal="center", vertical="center")
            c_date.fill = row_fill
            c_date.border = cell_border
            c_date.protection = no_lock

            # 6. مبلغ (تومان) - ذخیره به صورت عدد واقعی با فرمت حسابداری
            raw_amt = int(t.get("amount", 0))
            amt_val = raw_amt if is_income else -raw_amt
            c_amt = ws.cell(row=current_row, column=6, value=amt_val)
            c_amt.number_format = '+#,##0;-#,##0;0'
            c_amt.font = Font(name=font_family, size=11, bold=True, color="059669" if is_income else "DC2626")
            c_amt.alignment = ExcelGenerator._rtl_align(horizontal="center", vertical="center")
            c_amt.fill = row_fill
            c_amt.border = cell_border
            c_amt.protection = no_lock

            # 7. مربی مربوطه
            c_trainer = ws.cell(row=current_row, column=7, value=t.get("trainer_name", "-") or "-")
            c_trainer.font = font_regular
            c_trainer.alignment = ExcelGenerator._rtl_align(horizontal="center", vertical="center")
            c_trainer.fill = row_fill
            c_trainer.border = cell_border
            c_trainer.protection = no_lock

            # 8. توضیحات
            c_desc = ws.cell(row=current_row, column=8, value=t.get("description", "-") or "-")
            c_desc.font = font_regular
            c_desc.alignment = ExcelGenerator._rtl_align(horizontal="right", vertical="center", wrap_text=True, indent=1)
            c_desc.fill = row_fill
            c_desc.border = cell_border
            c_desc.protection = no_lock

            current_row += 1

        # Handle empty transactions case
        if not transactions:
            ws.merge_cells(f"A{current_row}:H{current_row}")
            empty_cell = ws[f"A{current_row}"]
            empty_cell.value = "هیچ تراکنشی در این بازه زمانی یافت نشد."
            empty_cell.font = Font(name=font_family, size=11, italic=True, color="64748B")
            empty_cell.alignment = ExcelGenerator._rtl_align(horizontal="center", vertical="center")
            empty_cell.fill = zebra_fill_white
            empty_cell.border = cell_border
            empty_cell.protection = no_lock
            ws.row_dimensions[current_row].height = 30
            current_row += 1

        # ----------------------------------------------------
        # Total Summary Row at Table Bottom
        # ----------------------------------------------------
        last_data_row = current_row - 1
        total_row = current_row
        ws.row_dimensions[total_row].height = 28

        ws.merge_cells(f"A{total_row}:E{total_row}")
        lbl_tot = ws[f"A{total_row}"]
        lbl_tot.value = "تراز و سود خالص نهایی تراکنش‌های جدول:"
        tot_lbl_font = Font(name=title_font_family, size=10.5, bold=True, color="1E293B")
        lbl_tot.font = tot_lbl_font
        lbl_tot.alignment = ExcelGenerator._rtl_align(horizontal="center", vertical="center")

        fill_total = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        border_top_thin_bottom_double = Border(
            top=border_thin,
            bottom=border_double,
            left=border_thin,
            right=border_thin
        )

        ExcelGenerator._style_range(ws, f"A{total_row}:E{total_row}", font=tot_lbl_font, fill=fill_total, border=border_top_thin_bottom_double, alignment=ExcelGenerator._rtl_align(horizontal="center", vertical="center"))

        # F: Formula for SUM of amounts
        val_tot = ws.cell(row=total_row, column=6)
        if transactions:
            val_tot.value = f"=SUM(F9:F{last_data_row})"
        else:
            val_tot.value = 0
        val_tot.number_format = '+#,##0;-#,##0;0'
        val_tot.font = Font(name=font_family, size=12, bold=True, color="047857" if profit >= 0 else "B91C1C")
        val_tot.alignment = ExcelGenerator._rtl_align(horizontal="center", vertical="center")
        val_tot.fill = fill_total
        val_tot.border = border_top_thin_bottom_double
        val_tot.protection = no_lock

        # G: Empty cell under مربی مربوطه (no redundant 'تومان')
        unit_tot = ws.cell(row=total_row, column=7, value="")
        unit_tot.font = Font(name=font_family, size=11, bold=True, color="64748B")
        unit_tot.alignment = ExcelGenerator._rtl_align(horizontal="center", vertical="center")
        unit_tot.fill = fill_total
        unit_tot.border = border_top_thin_bottom_double
        unit_tot.protection = no_lock

        # H: Spacer
        h_tot = ws.cell(row=total_row, column=8, value="")
        h_tot.alignment = ExcelGenerator._rtl_align(horizontal="center", vertical="center")
        h_tot.fill = fill_total
        h_tot.border = border_top_thin_bottom_double
        h_tot.protection = no_lock

        # ----------------------------------------------------
        # Column Widths Optimization (Auto-fit with safe minimums)
        # ----------------------------------------------------
        min_widths = {
            "A": 9,    # ردیف
            "B": 18,   # نوع تراکنش
            "C": 28,   # عنوان / طرف حساب
            "D": 20,   # دسته‌بندی
            "E": 15,   # تاریخ شمسی
            "F": 20,   # مبلغ (تومان)
            "G": 20,   # مربی مربوطه
            "H": 34,   # توضیحات
        }

        for col_idx in range(1, 9):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for r in range(8, total_row + 1):
                cell_v = ws.cell(row=r, column=col_idx).value
                if cell_v is not None:
                    s = str(cell_v)
                    if not s.startswith("="):
                        max_len = max(max_len, len(s))
            
            calculated_width = max_len + 5
            final_width = max(calculated_width, min_widths.get(col_letter, 15))
            ws.column_dimensions[col_letter].width = final_width

        wb.save(str(output_path))
        try:
            log_info(f"فایل اکسل صورت مالی با موفقیت ایجاد شد: {output_path}")
        except Exception:
            pass
        return str(output_path)
